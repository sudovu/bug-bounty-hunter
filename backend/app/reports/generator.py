"""Vulnerability and Scope Report Generator.

Creates professional, evidence-backed security reports in:
- Microsoft Word Document (.docx)
- Structured Plain Text (.txt)
- Markdown (.md)

Suitable for responsible disclosure submissions to authorized bug bounty programs.
"""
from datetime import datetime
from io import BytesIO
from typing import List, Optional
from sqlalchemy.orm import Session
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from backend.app.core.config import settings
from backend.app.db.models import Finding, Program, VulnerabilityReport
from backend.app.reports.remediation_catalog import get_detailed_solution


SEVERITY_COLORS = {
    "CRITICAL": RGBColor(220, 38, 38),   # Red
    "HIGH": RGBColor(234, 88, 12),       # Orange
    "MEDIUM": RGBColor(202, 138, 4),     # Amber/Yellow
    "LOW": RGBColor(37, 99, 235),        # Blue
    "INFO": RGBColor(100, 116, 139)      # Gray
}


def generate_finding_markdown_report(db: Session, finding: Finding, author: Optional[str] = None) -> str:
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    program = db.query(Program).filter(Program.id == finding.program_id).first()
    program_name = program.name if program else "Authorized Security Program"
    program_url = program.program_url if program else "N/A"
    detailed_solution = get_detailed_solution(finding.category, finding.title, finding.asset, finding.recommended_remediation)

    evidence_text = finding.evidence_summary
    if finding.evidence_items:
        evidence_text += "\n\n### Supporting Evidence Items:\n"
        for item in finding.evidence_items:
            evidence_text += f"- **Type:** `{item.evidence_type}`\n```\n{item.raw_evidence}\n```\n"

    report_md = f"""# VULNERABILITY REPORT: {finding.title}

---
**Report ID:** `{finding.unique_finding_id}`  
**Auditor / Reported By:** {author_name}  
**Program / Scope:** [{program_name}]({program_url})  
**Target Asset:** `{finding.asset}`  
**Severity:** **{finding.severity.value}**  
**Confidence:** {finding.confidence.value}  
**Status:** {finding.status.value}  
**Date Discovered:** {finding.discovered_at.strftime('%Y-%m-%d %H:%M:%S UTC')}  
**Detection Method:** `{finding.detection_method}`  
---

## 1. Executive Summary
During authorized security assessment on `{finding.asset}`, a security weakness categorized under **{finding.category}** was identified.
This finding has been assessed with **{finding.severity.value}** severity based on non-invasive technical evidence.

## 2. Technical Description
{finding.description}

- **Affected Component:** `{finding.affected_component or finding.asset}`
- **Security Standard / Policy Reference:** {finding.policy_reference or 'Responsible Disclosure Best Practices'}

## 3. Evidence & Verification
{evidence_text}

## 4. Potential Security Impact
{finding.potential_impact or 'Detailed impact depends on surrounding defensive architecture.'}

## 5. Safe Steps to Reproduce
The following non-destructive verification steps can safely confirm the reported behavior:

```bash
{finding.safe_reproduction or '# Safe curl check'}
```

> **Notice:** All testing was conducted strictly within authorized scope in compliance with the program policy. No exploits, payloads, or intrusive actions were performed.

## 6. Detailed Technical Solution & Remediation Plan
{detailed_solution}

---
*Report prepared by {author_name}. Bug Bounty Hunter Bot (Strict Scope & Ethical Auditing).*
"""
    return report_md


def generate_finding_text_report(db: Session, finding: Finding, author: Optional[str] = None) -> str:
    """Generate clean plain text (.txt) report suitable for text editors and form submissions."""
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    program = db.query(Program).filter(Program.id == finding.program_id).first()
    program_name = program.name if program else "Authorized Security Program"
    detailed_solution = get_detailed_solution(finding.category, finding.title, finding.asset, finding.recommended_remediation)

    evidence_text = finding.evidence_summary or "No raw evidence recorded."
    if finding.evidence_items:
        for item in finding.evidence_items:
            evidence_text += f"\n[{item.evidence_type}]\n{item.raw_evidence}\n"

    lines = [
        "=" * 76,
        f"VULNERABILITY REPORT: {finding.title.upper()}",
        "=" * 76,
        "",
        f"Report ID:           {finding.unique_finding_id}",
        f"Auditor / Reporter:  {author_name}",
        f"Program / Scope:     {program_name}",
        f"Target Asset:        {finding.asset}",
        f"Category:            {finding.category}",
        f"Severity:            {finding.severity.value}",
        f"Confidence:          {finding.confidence.value}",
        f"Date Discovered:     {finding.discovered_at.strftime('%Y-%m-%d %H:%M:%S UTC')}",
        f"Detection Method:    {finding.detection_method}",
        f"Status:              {finding.status.value}",
        "-" * 76,
        "",
        "1. EXECUTIVE SUMMARY",
        "-" * 30,
        f"During authorized assessment on {finding.asset}, a security issue was",
        f"identified and categorized as {finding.category} with {finding.severity.value} severity.",
        "",
        "2. TECHNICAL DESCRIPTION",
        "-" * 30,
        finding.description or "No description provided.",
        "",
        f"Affected Component:  {finding.affected_component or finding.asset}",
        f"Policy Reference:    {finding.policy_reference or 'Responsible Disclosure Best Practices'}",
        "",
        "3. TECHNICAL EVIDENCE & HEADERS",
        "-" * 30,
        evidence_text,
        "",
        "4. POTENTIAL SECURITY IMPACT",
        "-" * 30,
        finding.potential_impact or "Evaluation depends on surrounding infrastructure.",
        "",
        "5. SAFE REPRODUCTION STEPS",
        "-" * 30,
        finding.safe_reproduction or "# Safe curl inspection",
        "",
        "6. DETAILED TECHNICAL REMEDIATION PLAN",
        "-" * 30,
        detailed_solution,
        "",
        "=" * 76,
        f"Report prepared by: {author_name}",
        "Bug Bounty Hunter Bot (Strict Scope & Ethical Auditing)",
        "=" * 76
    ]
    return "\n".join(lines)


def generate_finding_docx_report(db: Session, finding: Finding, author: Optional[str] = None) -> bytes:
    """Generate professional Word Document (.docx) report with formatting and tables."""
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    program = db.query(Program).filter(Program.id == finding.program_id).first()
    program_name = program.name if program else "Authorized Security Program"
    detailed_solution = get_detailed_solution(finding.category, finding.title, finding.asset, finding.recommended_remediation)

    doc = Document()

    # Title Heading
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("VULNERABILITY ASSESSMENT REPORT")
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)  # Dark slate
    title_p.paragraph_format.space_after = Pt(4)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Finding: {finding.title}")
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105)
    sub_p.paragraph_format.space_after = Pt(14)

    # Metadata Table
    table = doc.add_table(rows=7, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    meta_rows = [
        ("Report Identifier", finding.unique_finding_id),
        ("Auditor / Reporter", author_name),
        ("Target Asset", finding.asset),
        ("Program / Scope", program_name),
        ("Severity Rating", finding.severity.value),
        ("Weakness Category", finding.category),
        ("Discovered Timestamp", finding.discovered_at.strftime('%Y-%m-%d %H:%M:%S UTC'))
    ]

    for i, (label, val) in enumerate(meta_rows):
        row = table.rows[i]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]

        lbl_run = cell_lbl.paragraphs[0].add_run(label)
        lbl_run.font.bold = True
        lbl_run.font.size = Pt(9.5)
        lbl_run.font.color.rgb = RGBColor(71, 85, 105)

        val_run = cell_val.paragraphs[0].add_run(val)
        val_run.font.size = Pt(9.5)

        if label == "Severity Rating":
            val_run.font.bold = True
            val_run.font.color.rgb = SEVERITY_COLORS.get(val.upper(), RGBColor(15, 23, 42))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for sections
    def add_section(heading: str, body: str, is_code: bool = False):
        h = doc.add_heading(level=2)
        h_run = h.add_run(heading)
        h_run.font.size = Pt(13)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(30, 41, 59)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)

        if is_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(body)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            p = doc.add_paragraph()
            run = p.add_run(body)
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(51, 65, 85)

    add_section("1. Executive Summary", 
        f"During authorized security auditing on {finding.asset}, a {finding.severity.value} severity "
        f"weakness was detected in {finding.category}. This finding was confirmed through passive/defensive verification.")

    add_section("2. Technical Description", finding.description or "No description provided.")

    evidence_text = finding.evidence_summary or "No raw evidence available."
    if finding.evidence_items:
        for item in finding.evidence_items:
            evidence_text += f"\n\n[{item.evidence_type}]\n{item.raw_evidence}"
    add_section("3. Evidence & Technical Observation", evidence_text, is_code=True)

    add_section("4. Security Impact Assessment", 
        finding.potential_impact or "The potential impact depends on surrounding perimeter defenses and access controls.")

    add_section("5. Non-Destructive Reproduction Steps", 
        finding.safe_reproduction or "# Safe curl command check", is_code=True)

    add_section("6. Detailed Technical Solution & Directives", detailed_solution)

    # Footer note
    doc.add_paragraph().paragraph_format.space_before = Pt(18)
    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run("This report was compiled in accordance with standard responsible disclosure practices.")
    footer_run.font.size = Pt(8.5)
    footer_run.font.italic = True
    footer_run.font.color.rgb = RGBColor(148, 163, 184)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_full_audit_docx_report(db: Session, program_id: Optional[int] = None, author: Optional[str] = None) -> bytes:
    """Generate a comprehensive multi-finding audit report in Word Document (.docx) format."""
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    query = db.query(Finding)
    if program_id:
        query = query.filter(Finding.program_id == program_id)
        prog = db.query(Program).filter(Program.id == program_id).first()
        prog_title = prog.name if prog else f"Program {program_id}"
    else:
        prog_title = "All Registered Programs"

    findings = query.order_by(Finding.discovered_at.desc()).all()

    doc = Document()

    # Title
    t_p = doc.add_paragraph()
    t_run = t_p.add_run("COMPREHENSIVE SECURITY AUDIT REPORT")
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Auditor: {author_name} | Scope: {prog_title} | Total Findings: {len(findings)} | Date: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_heading("Executive Summary & Findings Overview", level=2)

    if not findings:
        doc.add_paragraph("No vulnerabilities or weaknesses currently recorded for this scope.")
    else:
        # Overview Table
        table = doc.add_table(rows=len(findings) + 1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["ID", "Target", "Severity", "Finding Title"]
        for j, h in enumerate(headers):
            c_run = table.rows[0].cells[j].paragraphs[0].add_run(h)
            c_run.font.bold = True
            c_run.font.size = Pt(9.5)

        for i, f in enumerate(findings, 1):
            row = table.rows[i]
            row.cells[0].paragraphs[0].add_run(f.unique_finding_id).font.size = Pt(8.5)
            row.cells[1].paragraphs[0].add_run(f.asset).font.size = Pt(8.5)
            sev_run = row.cells[2].paragraphs[0].add_run(f.severity.value)
            sev_run.font.bold = True
            sev_run.font.size = Pt(8.5)
            sev_run.font.color.rgb = SEVERITY_COLORS.get(f.severity.value, RGBColor(15, 23, 42))
            row.cells[3].paragraphs[0].add_run(f.title).font.size = Pt(8.5)

        doc.add_page_break()

        # Detailed Findings Section
        for idx, f in enumerate(findings, 1):
            doc.add_heading(f"{idx}. {f.title}", level=2)
            p_meta = doc.add_paragraph()
            p_meta.add_run(f"Target: {f.asset} | Severity: {f.severity.value} | Category: {f.category}").font.bold = True

            doc.add_paragraph(f.description or "No description.")

            if f.evidence_summary:
                doc.add_heading("Evidence:", level=3)
                p_ev = doc.add_paragraph()
                p_ev.add_run(f.evidence_summary).font.name = "Consolas"
                p_ev.paragraph_format.left_indent = Inches(0.2)

            sol = get_detailed_solution(f.category, f.title, f.asset, f.recommended_remediation)
            doc.add_heading("Detailed Technical Solution & Remediation:", level=3)
            doc.add_paragraph(sol)

            doc.add_paragraph().paragraph_format.space_after = Pt(12)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def generate_full_audit_text_report(db: Session, program_id: Optional[int] = None, author: Optional[str] = None) -> str:
    """Generate a combined plain text (.txt) report for all findings."""
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    query = db.query(Finding)
    if program_id:
        query = query.filter(Finding.program_id == program_id)
        prog = db.query(Program).filter(Program.id == program_id).first()
        prog_title = prog.name if prog else f"Program {program_id}"
    else:
        prog_title = "All Registered Programs"

    findings = query.order_by(Finding.discovered_at.desc()).all()

    lines = [
        "=" * 80,
        f"COMPREHENSIVE SECURITY AUDIT REPORT: {prog_title.upper()}",
        f"Auditor / Prepared By: {author_name}",
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} | Total Findings: {len(findings)}",
        "=" * 80,
        "",
        "SUMMARY OF FINDINGS:",
        "-" * 80
    ]

    for idx, f in enumerate(findings, 1):
        lines.append(f"[{idx:02d}] [{f.severity.value:8s}] {f.title}")
        lines.append(f"     Target: {f.asset} | Category: {f.category} | ID: {f.unique_finding_id}")

    lines.append("")
    lines.append("=" * 80)
    lines.append("DETAILED TECHNICAL OBSERVATIONS & REMEDIATION")
    lines.append("=" * 80)

    for idx, f in enumerate(findings, 1):
        lines.append(f"\n[{idx}] {f.title.upper()} ({f.unique_finding_id})")
        lines.append(f"Target:      {f.asset}")
        lines.append(f"Severity:    {f.severity.value}")
        lines.append(f"Category:    {f.category}")
        lines.append(f"Description: {f.description or 'N/A'}")
        if f.evidence_summary:
            lines.append(f"Evidence:\n{f.evidence_summary}")
        sol = get_detailed_solution(f.category, f.title, f.asset, f.recommended_remediation)
        lines.append(f"Detailed Technical Remediation:\n{sol}")
        lines.append("-" * 40)

    return "\n".join(lines)


def generate_docx_from_markdown(markdown_text: str, title: str = "Vulnerability Report", author: Optional[str] = None) -> bytes:
    """Convert Markdown report text to a formatted Word Document (.docx)."""
    author_name = author or getattr(settings, "DEFAULT_REPORT_AUTHOR", "Bhuwan Gautam (IT Officer, Everest Wireless Network Pvt. Ltd.)")
    doc = Document()

    t_p = doc.add_paragraph()
    t_run = t_p.add_run(title)
    t_run.font.size = Pt(20)
    t_run.font.bold = True
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Auditor / Reporter: {author_name}")
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(71, 85, 105)

    in_code_block = False
    code_lines = []

    for line in markdown_text.splitlines():
        trimmed = line.strip()
        if trimmed.startswith("```"):
            if in_code_block:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(51, 65, 85)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if trimmed.startswith("# "):
            doc.add_heading(trimmed[2:], level=1)
        elif trimmed.startswith("## "):
            doc.add_heading(trimmed[3:], level=2)
        elif trimmed.startswith("### "):
            doc.add_heading(trimmed[4:], level=3)
        elif trimmed.startswith("- ") or trimmed.startswith("* "):
            doc.add_paragraph(trimmed[2:], style='List Bullet')
        elif trimmed in ("---", "***"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
        elif trimmed:
            doc.add_paragraph(line)

    if in_code_block and code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(9)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def create_and_save_report(db: Session, finding_id: int) -> VulnerabilityReport:
    finding = db.query(Finding).filter(Finding.id == finding_id).first()
    if not finding:
        raise ValueError(f"Finding {finding_id} not found")

    content = generate_finding_markdown_report(db, finding)
    report = VulnerabilityReport(
        program_id=finding.program_id,
        title=f"Report: {finding.title} ({finding.unique_finding_id})",
        format="MARKDOWN",
        content=content,
        created_at=datetime.utcnow()
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
