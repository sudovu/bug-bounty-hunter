"""Tests for Vulnerability Report Generation (Docx, Text, Markdown)."""
from datetime import datetime
from io import BytesIO
import pytest
from docx import Document
from fastapi.testclient import TestClient

from backend.app.main import app
from backend.app.db.base import get_db
from backend.app.db.models import (
    Program, Finding, FindingSeverity, FindingConfidence, FindingStatus, VulnerabilityReport
)
from backend.app.reports.generator import (
    generate_finding_text_report,
    generate_finding_docx_report,
    generate_full_audit_text_report,
    generate_full_audit_docx_report,
    generate_docx_from_markdown,
    create_and_save_report
)

client = TestClient(app)


@pytest.fixture
def sample_finding(db_session, sample_program):
    finding = Finding(
        unique_finding_id="BB-TEST-001",
        program_id=sample_program.id,
        asset="https://portal.example.com",
        category="Missing Security Headers",
        title="Missing Content-Security-Policy and HSTS Header",
        description="The target web server does not return Strict-Transport-Security or CSP headers.",
        evidence_summary="HTTP/1.1 200 OK\nServer: nginx\n(No HSTS header present)",
        affected_component="Web Front-End HTTPS",
        potential_impact="Users may be susceptible to clickjacking or SSL stripping on insecure connections.",
        safe_reproduction="curl -I https://portal.example.com",
        policy_reference="CWE-693: Protection Mechanism Failure",
        recommended_remediation="Implement Strict-Transport-Security: max-age=31536000; includeSubDomains and a restrictive CSP.",
        detection_method="safe_passive_check",
        severity=FindingSeverity.MEDIUM,
        confidence=FindingConfidence.HIGH,
        status=FindingStatus.DISCOVERED,
        deduplication_hash="hash_test_001",
        discovered_at=datetime.utcnow()
    )
    db_session.add(finding)
    db_session.commit()
    db_session.refresh(finding)
    return finding


def test_generate_finding_text_report(db_session, sample_finding):
    txt = generate_finding_text_report(db_session, sample_finding)
    assert isinstance(txt, str)
    assert sample_finding.unique_finding_id in txt
    assert "MISSING CONTENT-SECURITY-POLICY" in txt.upper()
    assert sample_finding.asset in txt
    assert "REMEDIATION RECOMMENDATIONS" in txt


def test_generate_finding_docx_report(db_session, sample_finding):
    docx_bytes = generate_finding_docx_report(db_session, sample_finding)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000

    # Ensure python-docx can parse the generated file
    bio = BytesIO(docx_bytes)
    doc = Document(bio)
    paragraphs_text = " ".join(p.text for p in doc.paragraphs)
    assert "VULNERABILITY ASSESSMENT REPORT" in paragraphs_text
    assert sample_finding.title in paragraphs_text
    assert sample_finding.asset in paragraphs_text


def test_generate_full_audit_reports(db_session, sample_finding):
    txt = generate_full_audit_text_report(db_session, sample_finding.program_id)
    assert "COMPREHENSIVE SECURITY AUDIT REPORT" in txt
    assert sample_finding.unique_finding_id in txt

    docx_bytes = generate_full_audit_docx_report(db_session, sample_finding.program_id)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000

    doc = Document(BytesIO(docx_bytes))
    paragraphs_text = " ".join(p.text for p in doc.paragraphs)
    assert "COMPREHENSIVE SECURITY AUDIT REPORT" in paragraphs_text


def test_generate_docx_from_markdown():
    sample_md = """# Security Audit
## Overview
This is an overview paragraph.

- Item 1
- Item 2

```bash
curl -I https://example.com
```
"""
    docx_bytes = generate_docx_from_markdown(sample_md, title="Test Report")
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000

    doc = Document(BytesIO(docx_bytes))
    full_text = " ".join(p.text for p in doc.paragraphs)
    assert "Test Report" in full_text
    assert "Overview" in full_text
    assert "Item 1" in full_text


def test_api_export_finding_docx_and_txt(db_session, sample_finding):
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        # Export DOCX
        res = client.get(f"/api/v1/reports/finding/{sample_finding.id}/export?format=docx")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert len(res.content) > 1000

        # Export TXT
        res_txt = client.get(f"/api/v1/reports/finding/{sample_finding.id}/export?format=txt")
        assert res_txt.status_code == 200
        assert "text/plain" in res_txt.headers["content-type"]
        assert sample_finding.unique_finding_id in res_txt.text

        # Export MD
        res_md = client.get(f"/api/v1/reports/finding/{sample_finding.id}/export?format=md")
        assert res_md.status_code == 200
        assert "text/markdown" in res_md.headers["content-type"]
        assert sample_finding.title in res_md.text
    finally:
        app.dependency_overrides.clear()


def test_api_export_all(db_session, sample_finding):
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        res_docx = client.get("/api/v1/reports/export-all?format=docx")
        assert res_docx.status_code == 200
        assert res_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        res_txt = client.get("/api/v1/reports/export-all?format=txt")
        assert res_txt.status_code == 200
        assert "text/plain" in res_txt.headers["content-type"]
        assert sample_finding.title in res_txt.text
    finally:
        app.dependency_overrides.clear()


def test_api_download_saved_report(db_session, sample_finding):
    app.dependency_overrides[get_db] = lambda: db_session
    try:
        saved_rep = create_and_save_report(db_session, sample_finding.id)
        
        # Download DOCX
        res_docx = client.get(f"/api/v1/reports/{saved_rep.id}/download?format=docx")
        assert res_docx.status_code == 200
        assert res_docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        # Download TXT
        res_txt = client.get(f"/api/v1/reports/{saved_rep.id}/download?format=txt")
        assert res_txt.status_code == 200
        assert "text/plain" in res_txt.headers["content-type"]

        # Download MD
        res_md = client.get(f"/api/v1/reports/{saved_rep.id}/download?format=md")
        assert res_md.status_code == 200
        assert "text/markdown" in res_md.headers["content-type"]
    finally:
        app.dependency_overrides.clear()
